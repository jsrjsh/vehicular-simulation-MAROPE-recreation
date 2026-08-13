# /// script
# requires-python = ">=3.11"
# dependencies = [
#     "python-docx",
# ]
# ///

from docx import Document
import sys

def generate_inventory(docx_path):
    doc = Document(docx_path)
    
    headings = []
    paragraphs_per_sec = []
    curr_para_count = 0
    tables_count = len(doc.tables)
    figures_count = 0
    equations_count = 0
    algorithms_count = 0
    citations_count = 0
    bib_count = 0
    footnotes_count = 0
    
    for para in doc.paragraphs:
        if para.style.name.startswith('Heading'):
            if headings:
                paragraphs_per_sec.append(curr_para_count)
            headings.append(para.text)
            curr_para_count = 0
        else:
            curr_para_count += 1
            text = para.text.strip()
            if 'Fig.' in text or 'Figure' in text:
                if len(text) < 200:
                    figures_count += 1
            if 'Algorithm' in text:
                if len(text) < 100:
                    algorithms_count += 1
            if 'Eq.' in text or 'Equation' in text:
                pass
            if '[' in text and ']' in text:
                citations_count += text.count('[')
            
            # Simple heuristic for equations (they often contain =)
            if '=' in text and len(text) < 100 and not text.endswith('.'):
                equations_count += 1
                
    if headings:
        paragraphs_per_sec.append(curr_para_count)
        
    with open('DOCX_CONTENT_INVENTORY.md', 'w') as f:
        f.write("# Manuscript Inventory\n\n")
        f.write("## Headings and Paragraphs\n")
        for i, h in enumerate(headings):
            f.write(f"- {h} ({paragraphs_per_sec[i]} paragraphs)\n")
            
        f.write(f"\n## Objects\n")
        f.write(f"- Tables: {tables_count}\n")
        f.write(f"- Figures (Estimated): {figures_count}\n")
        f.write(f"- Algorithms (Estimated): {algorithms_count}\n")
        f.write(f"- Equations (Estimated): {equations_count}\n")
        f.write(f"- Citations (Estimated): {citations_count}\n")
        f.write(f"- Bibliography Entries: {bib_count}\n")
        f.write(f"- Footnotes: {footnotes_count}\n")
        
        f.write("\n## Mapping Table\n")
        f.write("| Object | Destination LaTeX File | Status |\n")
        f.write("|---|---|---|\n")
        for i, h in enumerate(headings):
            f.write(f"| Section: {h} | sections/sec_{i:02d}.tex | PENDING |\n")

if __name__ == "__main__":
    generate_inventory("manuscript/submission/MA_ROPE_IEEE_Paper_v21.docx")
