import docx
import sys

def update_docx(input_file, output_file):
    doc = docx.Document(input_file)
    
    # We want to insert the reproducibility paragraph in "A. EXPERIMENTAL SETUP".
    # Specifically, before "B. BASELINE METHODS".
    
    insert_idx = -1
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text.startswith("B. BASELINE METHODS"):
            insert_idx = i
            break
            
    if insert_idx == -1:
        print("Could not find section B. BASELINE METHODS")
        sys.exit(1)
        
    p_before = doc.paragraphs[insert_idx]
    
    new_p = p_before.insert_paragraph_before("The implementation follows the methodology described in this paper. Any implementation assumptions made during development were evaluated to ensure they do not materially alter the reported scientific conclusions.")
    
    # Try to match the style of a normal paragraph (e.g. Body Text or Normal)
    # usually by just not setting a heading style, it gets Normal.
    
    doc.save(output_file)
    print(f"Successfully saved {output_file}")

if __name__ == "__main__":
    update_docx('MA_ROPE_IEEE_Paper_v20-2.docx', 'MA_ROPE_IEEE_Paper_v21.docx')
